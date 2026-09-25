import sys, os, json, sqlite3, urllib.request, tempfile, time, shutil, re, html
from pathlib import Path
from PySide6.QtCore import Qt, Signal, QUrl, QSize
from PySide6.QtGui import QFont, QColor, QPixmap, QPainter, QPen, QTextCharFormat, QTextCursor, QTextImageFormat
from PySide6.QtWidgets import *
from PySide6.QtWebEngineWidgets import QWebEngineView
from PySide6.QtPdf import QPdfDocument
from PySide6.QtPdfWidgets import QPdfView
from docx import Document
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from coverai_studio import CoverAIStudio

DATA_DIR = Path.home() / "CuadernoJuridicoData"
DATA_DIR.mkdir(exist_ok=True)
DB_PATH = DATA_DIR / "datos.db"
CFG_PATH = DATA_DIR / "config.json"
DEFAULT_CFG = {"app_name":"Mi Cuaderno Jurídico","user_name":"Kelvin","university":"Universidad","career":"Ciencias Jurídicas","accent":"#D4AF37","theme":"dark"}

def load_cfg():
    try:
        return {**DEFAULT_CFG, **json.loads(CFG_PATH.read_text(encoding="utf-8"))}
    except Exception:
        return DEFAULT_CFG.copy()

def save_cfg(cfg): CFG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")

class Store:
    def __init__(self):
        self.db=sqlite3.connect(DB_PATH); self.db.row_factory=sqlite3.Row
        self.db.executescript("""
        CREATE TABLE IF NOT EXISTS subjects(id INTEGER PRIMARY KEY,name TEXT NOT NULL,icon TEXT DEFAULT '⚖');
        CREATE TABLE IF NOT EXISTS units(id INTEGER PRIMARY KEY,subject_id INTEGER,title TEXT NOT NULL,content TEXT DEFAULT '',completed INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS cards(id INTEGER PRIMARY KEY,subject_id INTEGER,front TEXT NOT NULL,back TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS library(id INTEGER PRIMARY KEY,name TEXT,path TEXT,type TEXT);
        CREATE TABLE IF NOT EXISTS covers(subject_id INTEGER PRIMARY KEY,title TEXT,subtitle TEXT,author TEXT,university TEXT,faculty TEXT,career TEXT,year TEXT,color TEXT,icon TEXT,university_logo TEXT,faculty_logo TEXT,background_image TEXT,show_university_logo INTEGER DEFAULT 1,show_faculty_logo INTEGER DEFAULT 1,show_author INTEGER DEFAULT 1,show_university INTEGER DEFAULT 1,show_career INTEGER DEFAULT 1,show_year INTEGER DEFAULT 1);
        CREATE TABLE IF NOT EXISTS global_cover(id INTEGER PRIMARY KEY CHECK(id=1),title TEXT,subtitle TEXT,author TEXT,university TEXT,faculty TEXT,career TEXT,year TEXT,color TEXT,icon TEXT,university_logo TEXT,faculty_logo TEXT,background_image TEXT,background_position TEXT DEFAULT 'center',font_family TEXT DEFAULT 'Georgia',title_font_size INTEGER DEFAULT 30,text_font_size INTEGER DEFAULT 16,show_university_logo INTEGER DEFAULT 1,show_faculty_logo INTEGER DEFAULT 1,show_author INTEGER DEFAULT 1,show_university INTEGER DEFAULT 1,show_career INTEGER DEFAULT 1,show_year INTEGER DEFAULT 1);
        """)
        cols=[r[1] for r in self.db.execute("PRAGMA table_info(units)").fetchall()]
        if 'docx_path' not in cols: self.db.execute("ALTER TABLE units ADD COLUMN docx_path TEXT DEFAULT ''"); self.db.commit()
        cover_cols=[r[1] for r in self.db.execute("PRAGMA table_info(global_cover)").fetchall()]
        for name,definition in [('background_position',"TEXT DEFAULT 'center'"),('font_family',"TEXT DEFAULT 'Georgia'"),('font_color',"TEXT DEFAULT '#FFFFFF'"),('title_font_size',"INTEGER DEFAULT 30"),('text_font_size',"INTEGER DEFAULT 16"),('frame_style',"TEXT DEFAULT 'Doble dorado'"),('margin_size',"INTEGER DEFAULT 36")]:
            if name not in cover_cols: self.db.execute(f"ALTER TABLE global_cover ADD COLUMN {name} {definition}")
        for table in ('global_cover','covers'):
            current_cols=[r[1] for r in self.db.execute(f"PRAGMA table_info({table})").fetchall()]
            if 'coverai_path' not in current_cols:self.db.execute(f"ALTER TABLE {table} ADD COLUMN coverai_path TEXT DEFAULT ''")
        self.db.commit()
        if self.value("SELECT COUNT(*) FROM subjects")==0: self.seed()
    def rows(self,q,p=()): return self.db.execute(q,p).fetchall()
    def row(self,q,p=()): return self.db.execute(q,p).fetchone()
    def value(self,q,p=()): return self.db.execute(q,p).fetchone()[0]
    def run(self,q,p=()): self.db.execute(q,p); self.db.commit()
    def seed(self):
        cur=self.db.execute("INSERT INTO subjects(name,icon) VALUES(?,?)",("Teoría General del Proceso","⚖")); sid=cur.lastrowid
        self.db.executemany("INSERT INTO subjects(name,icon) VALUES(?,?)",[("Derecho Constitucional","🏛"),("Derecho Civil","📜")])
        names=["Derecho procesal","Jurisdicción","Competencia","Derecho de acción","Pretensión","Derecho de excepción","El proceso","Relación jurídica procesal","Resoluciones judiciales","Recursos","Cosa juzgada"]
        self.db.executemany("INSERT INTO units(subject_id,title,content,completed) VALUES(?,?,?,?)",[(sid,f"Unidad {i+1}: {n}","Escribe aquí tus conceptos, artículos, ejemplos y conclusiones.",1 if i==0 else 0) for i,n in enumerate(names)])
        self.db.executemany("INSERT INTO cards(subject_id,front,back) VALUES(?,?,?)",[(sid,"¿Qué es la jurisdicción?","La función estatal destinada a resolver conflictos jurídicos."),(sid,"¿Qué es la competencia?","El ámbito dentro del cual un órgano jurisdiccional ejerce su función."),(sid,"¿Qué es la cosa juzgada?","La autoridad y eficacia de una resolución firme.")]); self.db.commit()

class TextDialog(QDialog):
    def __init__(self,title,fields,parent=None):
        super().__init__(parent); self.setWindowTitle(title); self.setMinimumWidth(420); lay=QFormLayout(self); self.ed={}
        for key,label,multi in fields:
            w=QTextEdit() if multi else QLineEdit();
            if multi: w.setFixedHeight(100)
            self.ed[key]=w; lay.addRow(label,w)
        box=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel); box.accepted.connect(self.accept); box.rejected.connect(self.reject); lay.addRow(box)
    def values(self): return {k:(w.toPlainText() if isinstance(w,QTextEdit) else w.text()).strip() for k,w in self.ed.items()}

class CoverAppearanceDialog(QDialog):
    FONTS=['Georgia','Times New Roman','Garamond','Arial','Calibri','Verdana']
    FRAMES=['Ninguno','Doble dorado','Expediente clásico','Romano institucional','Borgoña jurídico','Azul notarial','Minimalista negro','Clásico académico','Institucional elegante']
    def __init__(self,data,parent=None):
        super().__init__(parent); self.setWindowTitle('Estilo de la portada'); self.setMinimumWidth(430); form=QFormLayout(self)
        self.font=QComboBox(); self.font.addItems(self.FONTS); self.font.setCurrentText(data.get('font_family','Georgia'))
        self.font_color_value=data.get('font_color','#FFFFFF') or '#FFFFFF'; self.font_color=QPushButton('Seleccionar color del texto'); self.font_color.clicked.connect(self.choose_font_color); self.update_color_button()
        self.frame=QComboBox(); self.frame.addItems(self.FRAMES); self.frame.setCurrentText(data.get('frame_style','Doble dorado'))
        note=QLabel('La portada se compone automáticamente: todo centrado, tamaños equilibrados y ajuste para evitar recortes.'); note.setWordWrap(True); note.setObjectName('muted')
        form.addRow('Tipo de letra',self.font); form.addRow('Color del texto',self.font_color); form.addRow('Marco jurídico',self.frame); form.addRow(note)
        buttons=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); form.addRow(buttons)
    def update_color_button(self):self.font_color.setStyleSheet(f'QPushButton{{background:{self.font_color_value};color:#111827;border:1px solid #64748B;border-radius:6px;padding:7px}}')
    def choose_font_color(self):
        color=QColorDialog.getColor(QColor(self.font_color_value),self,'Color del texto')
        if color.isValid():self.font_color_value=color.name().upper();self.update_color_button()
    def values(self):return dict(font_family=self.font.currentText(),font_color=self.font_color_value,frame_style=self.frame.currentText())
class CoverAssetsDialog(QDialog):
    def __init__(self,data,parent=None):
        super().__init__(parent); self.setWindowTitle('Logos e imagen de fondo'); self.setMinimumWidth(520); self.uni_logo=data.get('university_logo','') or ''; self.fac_logo=data.get('faculty_logo','') or ''; self.bg_image=data.get('background_image','') or ''
        form=QFormLayout(self); self.fields=[]
        for label,kind,attr in [('Logo universidad','uni','uni_logo'),('Logo facultad','fac','fac_logo'),('Imagen de fondo','bg','bg_image')]:
            field=QLineEdit(getattr(self,attr)); field.setReadOnly(True); select=QPushButton('Seleccionar'); select.clicked.connect(lambda _,k=kind:self.pick(k)); clear=QPushButton('Eliminar'); clear.clicked.connect(lambda _,k=kind:self.clear(k)); row=QHBoxLayout(); row.addWidget(field,1); row.addWidget(select); row.addWidget(clear); host=QWidget(); host.setLayout(row); form.addRow(label,host); self.fields.append((kind,field))
        self.background_color=QComboBox(); [self.background_color.addItem(n,c) for n,c in MainCoverDialog.COLORS]; self.background_color.setCurrentIndex(max(0,self.background_color.findData(data.get('color','#146C82')))); form.addRow('Color de fondo',self.background_color)
        self.position=QComboBox(); self.position.addItems(['Centro','Arriba','Abajo','Izquierda','Derecha']); self.position.setCurrentText({'center':'Centro','top':'Arriba','bottom':'Abajo','left':'Izquierda','right':'Derecha'}.get(data.get('background_position','center'),'Centro')); form.addRow('Posición del fondo',self.position)
        note=QLabel('Si existe una imagen de fondo, se mostrará la imagen. Si se elimina, se mostrará el color de portada.'); note.setWordWrap(True); form.addRow(note)
        buttons=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); form.addRow(buttons)
    def pick(self,kind):
        path,_=QFileDialog.getOpenFileName(self,'Seleccionar imagen','','Imágenes (*.png *.jpg *.jpeg *.bmp *.webp);;Todos (*.*)')
        if not path:return
        if kind=='uni':self.uni_logo=path
        elif kind=='fac':self.fac_logo=path
        else:self.bg_image=path
        self.refresh_fields()
    def clear(self,kind):
        if kind=='uni':self.uni_logo=''
        elif kind=='fac':self.fac_logo=''
        else:self.bg_image=''
        self.refresh_fields()
    def refresh_fields(self):
        values={'uni':self.uni_logo,'fac':self.fac_logo,'bg':self.bg_image}
        for kind,field in self.fields:field.setText(values[kind])
    def values(self):
        positions={'Centro':'center','Arriba':'top','Abajo':'bottom','Izquierda':'left','Derecha':'right'}
        return dict(university_logo=self.uni_logo,faculty_logo=self.fac_logo,background_image=self.bg_image,color=self.background_color.currentData(),background_position=positions[self.position.currentText()])

class EditorCoverPreview(QLabel):
    def __init__(self,parent=None):
        super().__init__(parent); self.cover_color='#146C82'; self.cover_pixmap=QPixmap(); self.cover_position='center'; self.frame_style='Doble dorado'; self.setAlignment(Qt.AlignCenter); self.setWordWrap(True)
    def set_cover_background(self,color,path='',position='center'):
        self.cover_color=color or '#146C82'; self.cover_pixmap=QPixmap(path) if path and Path(path).exists() else QPixmap(); self.cover_position=position or 'center'; self.update()
    def set_frame_style(self,style): self.frame_style=style or 'Doble dorado'; self.update()
    def paintEvent(self,event):
        painter=QPainter(self); rect=self.rect(); painter.fillRect(rect,QColor(self.cover_color))
        if not self.cover_pixmap.isNull() and rect.width()>0 and rect.height()>0:
            image=self.cover_pixmap.scaled(rect.size(),Qt.KeepAspectRatioByExpanding,Qt.SmoothTransformation); x=(rect.width()-image.width())//2; y=(rect.height()-image.height())//2
            if self.cover_position=='top': y=0
            elif self.cover_position=='bottom': y=rect.height()-image.height()
            elif self.cover_position=='left': x=0
            elif self.cover_position=='right': x=rect.width()-image.width()
            painter.drawPixmap(x,y,image)
        painter.end(); super().paintEvent(event); painter=QPainter(self); paint_cover_frame(painter,self.rect().adjusted(2,2,-2,-2),self.frame_style); painter.end()

FRAME_STYLES={
    'Ninguno':('0px none transparent','transparent'),
    'Doble dorado':('7px double #D9A06C','#D9A06C'),
    'Expediente clásico':('10px ridge #8B6B3F','#8B6B3F'),
    'Romano institucional':('8px double #D7EEF5','#D7EEF5'),
    'Borgoña jurídico':('8px solid #B86B77','#B86B77'),
    'Azul notarial':('8px groove #4EA1D3','#4EA1D3'),
    'Minimalista negro':('4px solid #111827','#111827'),
}

def paint_cover_frame(painter,rect,style):
    if style=='Ninguno': return
    border,color=FRAME_STYLES.get(style,FRAME_STYLES['Doble dorado']); width=int(border.split('px')[0]); painter.setBrush(Qt.NoBrush)
    if style=='Doble dorado':
        painter.setPen(QPen(QColor(color),max(2,width//2),Qt.SolidLine)); painter.drawRect(rect.adjusted(width//2,width//2,-width//2,-width//2)); painter.setPen(QPen(QColor(color),max(1,width//3),Qt.SolidLine)); painter.drawRect(rect.adjusted(width,width,-width,-width))
    else:
        painter.setPen(QPen(QColor(color),width,Qt.SolidLine)); painter.drawRect(rect.adjusted(width//2,width//2,-width//2,-width//2))

def build_cover_html(data):
    esc=lambda value:html.escape(str(value or '')); details=[]
    if data.get('show_author',1) and data.get('author'):details.append(esc(data['author']))
    if data.get('faculty'):details.append(esc(data['faculty']))
    if data.get('show_career',1) and data.get('career'):details.append(esc(data['career']))
    if data.get('show_university',1) and data.get('university'):details.append(esc(data['university']))
    if data.get('show_year',1) and data.get('year'):details.append(esc(data['year']))
    def img(path,show):
        if not show or not path or not Path(path).exists():return ''
        return f"<img src='{QUrl.fromLocalFile(path).toString()}' width='54'><br>"
    font=html.escape(data.get('font_family') or 'Georgia'); color=html.escape(data.get('font_color') or '#FFFFFF')
    title_size=int(data.get('title_font_size') or 30); text_size=int(data.get('text_font_size') or 16)
    title=esc(data.get('title') or 'TÍTULO DEL CUADERNO').upper(); subtitle=esc(data.get('subtitle')).upper(); icon=esc(data.get('icon') or '⚖')
    return f"<div align='center' style='color:{color};font-family:{font};'><div>{img(data.get('university_logo',''),data.get('show_university_logo',1))}</div><div style='font-size:{text_size}px;letter-spacing:1px'>{subtitle}</div><div style='font-size:32px;margin:8px'>{icon}</div><div style='font-size:{max(18,title_size-9)}px;font-weight:700;line-height:1.08'>{title}</div><div style='margin-top:16px;font-size:{text_size}px;line-height:1.25'>{'<br>'.join(details)}</div><div style='margin-top:12px'>{img(data.get('faculty_logo',''),data.get('show_faculty_logo',1))}</div></div>"

def cover_pdf_font():
    candidates=[Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/'georgia.ttf',Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/'times.ttf',Path('/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf')]
    for path in candidates:
        try:
            if path.exists():
                if 'CoverPDF' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('CoverPDF',str(path)))
                return 'CoverPDF'
        except Exception:pass
    return 'Times-Roman'
def cover_wrapped_lines(text,font,size,max_width):
    words=str(text or '').split();lines=[];current=''
    for word in words:
        test=(current+' '+word).strip()
        if pdfmetrics.stringWidth(test,font,size)<=max_width:current=test
        else:
            if current:lines.append(current)
            current=word
    if current:lines.append(current)
    return lines or ['']
def cover_image(c,path,page_w,page_h,position='center'):
    if not path or not Path(path).exists():return False
    try:
        img=ImageReader(path);iw,ih=img.getSize();scale=max(page_w/iw,page_h/ih);dw,dh=iw*scale,ih*scale;x=(page_w-dw)/2;y=(page_h-dh)/2
        if position=='top':y=page_h-dh
        elif position=='bottom':y=0
        elif position=='left':x=0
        elif position=='right':x=page_w-dw
        c.drawImage(img,x,y,dw,dh,mask='auto');return True
    except Exception:return False
def cover_logo(c,path,x,y,max_size):
    if not path or not Path(path).exists():return y
    try:
        img=ImageReader(path);iw,ih=img.getSize();scale=min(max_size/iw,max_size/ih);w,h=iw*scale,ih*scale;c.drawImage(img,x-w/2,y-h,w,h,mask='auto');return y-h-8
    except Exception:return y
def create_cover_pdf_file(data,path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);w,h=letter;c=pdf_canvas.Canvas(str(path),pagesize=letter)
    if not cover_image(c,data.get('background_image',''),w,h,data.get('background_position','center')):
        c.setFillColor(HexColor(data.get('color') or '#146C82'));c.rect(0,0,w,h,fill=1,stroke=0)
    frame=data.get('frame_style','Doble dorado');frame_colors={'Doble dorado':'#D9A06C','Expediente clásico':'#8B6B3F','Romano institucional':'#D7EEF5','Borgoña jurídico':'#B86B77','Azul notarial':'#4EA1D3','Minimalista negro':'#111827','Clásico académico':'#C9A227','Institucional elegante':'#E7EDF3'}
    if frame!='Ninguno':
        fc=HexColor(frame_colors.get(frame,'#D9A06C'));c.setStrokeColor(fc);c.setLineWidth(3);c.rect(18,18,w-36,h-36)
        if frame=='Doble dorado':c.setLineWidth(1.4);c.rect(25,25,w-50,h-50)
    font=cover_pdf_font();color=data.get('font_color') or '#FFFFFF';margin=max(28,int(data.get('margin_size') or 36));usable=w-(margin*2)-30;x=w/2
    title_size=max(16,int(data.get('title_font_size') or 30));text_size=max(9,int(data.get('text_font_size') or 16));title_lines=cover_wrapped_lines(str(data.get('title') or 'TÍTULO DEL CUADERNO').upper(),font,title_size,usable)
    details=[]
    if data.get('show_author',1) and data.get('author'):details.append(str(data['author']))
    if data.get('faculty'):details.append(str(data['faculty']))
    if data.get('show_career',1) and data.get('career'):details.append(str(data['career']))
    if data.get('show_university',1) and data.get('university'):details.append(str(data['university']))
    if data.get('show_year',1) and data.get('year'):details.append(str(data['year']))
    # Fit everything vertically by reducing title and body sizes together.
    logo_top=72 if data.get('show_university_logo',1) and data.get('university_logo') else 0;logo_bottom=62 if data.get('show_faculty_logo',1) and data.get('faculty_logo') else 0
    while True:
        title_lines=cover_wrapped_lines(str(data.get('title') or 'TÍTULO DEL CUADERNO').upper(),font,title_size,usable)
        required=logo_top+30+40+len(title_lines)*(title_size*1.18)+40+sum(max(1,len(cover_wrapped_lines(d,font,text_size,usable)))*(text_size*1.2)+7 for d in details)+logo_bottom+70
        if required<=h-(margin*2) or (title_size<=16 and text_size<=8):break
        if title_size>16:title_size-=1
        if text_size>8:text_size-=1
    y=h-margin
    if data.get('show_university_logo',1):y=cover_logo(c,data.get('university_logo',''),x,y,72)
    c.setFillColor(HexColor(color));c.setFont(font,max(8,text_size-1));c.drawCentredString(x,y,str(data.get('subtitle') or '').upper());y-=35
    # Safe legal ornament; exact on preview and PDF because preview displays this PDF.
    c.setFont(font,18);c.drawCentredString(x,y,'*');y-=38
    c.setFont(font,title_size)
    for line in title_lines:c.drawCentredString(x,y,line);y-=title_size*1.18
    y-=30;c.setFont(font,text_size)
    for detail in details:
        for line in cover_wrapped_lines(detail,font,text_size,usable):c.drawCentredString(x,y,line);y-=text_size*1.2
        y-=7
    if data.get('show_faculty_logo',1):cover_logo(c,data.get('faculty_logo',''),x,max(margin+72,y-8),58)
    c.showPage();c.save();return path

class MainCoverDialog(QDialog):
    COLORS=[('Azul jurídico','#146C82'),('Azul institucional','#173B57'),('Borgoña universitario','#7A263A'),('Verde académico','#1F6F5F'),('Morado sobrio','#514078'),('Negro y dorado','#171A21')]
    def __init__(self,data,parent=None):
        super().__init__(parent); self.setWindowTitle('Portada principal'); self.resize(790,540); self.setMinimumSize(700,500); self.setWindowFlag(Qt.WindowMaximizeButtonHint,True); self.data=dict(data); self.uni_logo=self.data.get('university_logo','') or ''; self.fac_logo=self.data.get('faculty_logo','') or ''; self.bg_image=self.data.get('background_image','') or ''; self.background_position=self.data.get('background_position','center'); self.font_family=self.data.get('font_family','Georgia'); self.font_color=self.data.get('font_color','#FFFFFF'); self.frame_style=self.data.get('frame_style','Doble dorado'); self.margin_size=int(self.data.get('margin_size',36)); self.title_font_size=int(self.data.get('title_font_size',30)); self.text_font_size=int(self.data.get('text_font_size',16)); self.word_path=self.data.get('coverai_path','') or ''; self.layout_settings={}
        root=QHBoxLayout(self); root.setContentsMargins(5,5,5,5); root.setSpacing(6); controls=QFrame(); controls.setObjectName('card'); form=QFormLayout(controls); form.setContentsMargins(8,8,8,8); form.setHorizontalSpacing(7); form.setVerticalSpacing(3)
        self.title=QLineEdit(self.data.get('title','')); self.subtitle=QLineEdit(self.data.get('subtitle','Cuaderno de estudio')); self.author=QLineEdit(self.data.get('author','')); self.university=QLineEdit(self.data.get('university','')); self.faculty=QLineEdit(self.data.get('faculty','')); self.career=QLineEdit(self.data.get('career','')); self.year=QLineEdit(self.data.get('year',str(time.localtime().tm_year)))
        self.palette=QComboBox(); [self.palette.addItem(n,c) for n,c in self.COLORS]; idx=self.palette.findData(self.data.get('color','#146C82')); self.palette.setCurrentIndex(max(0,idx))
        for label,w in [('Título',self.title),('Subtítulo',self.subtitle),('Autor',self.author),('Universidad',self.university),('Facultad',self.faculty),('Carrera',self.career),('Año',self.year)]: form.addRow(label,w)
        appearance=QPushButton('Estilo de portada'); appearance.clicked.connect(self.edit_appearance); form.addRow('Diseño',appearance)
        assets=QPushButton('Logos e imagen de fondo'); assets.clicked.connect(self.edit_assets); form.addRow('Recursos',assets); word_btn=QPushButton('✎ Diseñar portada en Word'); word_btn.clicked.connect(self.design_in_word); form.addRow('Edición avanzada',word_btn)
        self.show_uni_logo=QCheckBox('Mostrar logo universidad'); self.show_uni_logo.setChecked(bool(self.data.get('show_university_logo',1))); self.show_fac_logo=QCheckBox('Mostrar logo facultad'); self.show_fac_logo.setChecked(bool(self.data.get('show_faculty_logo',1))); self.show_author=QCheckBox('Mostrar autor'); self.show_author.setChecked(bool(self.data.get('show_author',1))); self.show_university=QCheckBox('Mostrar universidad'); self.show_university.setChecked(bool(self.data.get('show_university',1))); self.show_career=QCheckBox('Mostrar carrera'); self.show_career.setChecked(bool(self.data.get('show_career',1))); self.show_year=QCheckBox('Mostrar año'); self.show_year.setChecked(bool(self.data.get('show_year',1)))
        for check in (self.show_uni_logo,self.show_fac_logo,self.show_author,self.show_university,self.show_career,self.show_year): check.setStyleSheet('QCheckBox{color:#FFFFFF;spacing:8px} QCheckBox::indicator{width:16px;height:16px;background:#FFFFFF;border:1px solid #94A3B8;border-radius:3px} QCheckBox::indicator:checked{background:#111827;border:1px solid #000000}')
        checks=QGridLayout(); checks.addWidget(self.show_uni_logo,0,0); checks.addWidget(self.show_fac_logo,0,1); checks.addWidget(self.show_author,1,0); checks.addWidget(self.show_university,1,1); checks.addWidget(self.show_career,2,0); checks.addWidget(self.show_year,2,1); wc=QWidget(); wc.setLayout(checks); form.addRow('Elementos visibles',wc)
        buttons=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); form.addRow(buttons); root.addWidget(controls,1)
        preview_card=QFrame(); preview_card.setObjectName('card'); pv=QVBoxLayout(preview_card); label=QLabel('Vista previa exacta tamaño Carta'); label.setObjectName('heading'); pv.addWidget(label); self.preview_pdf_doc=QPdfDocument(self); self.preview=QPdfView(); self.preview.setDocument(self.preview_pdf_doc); self.preview.setPageMode(QPdfView.PageMode.SinglePage); self.preview.setZoomMode(QPdfView.ZoomMode.FitInView); self.preview.setMinimumSize(330,430); pv.addWidget(self.preview,1); root.addWidget(preview_card,1)
        widgets=[self.title,self.subtitle,self.author,self.university,self.faculty,self.career,self.year ]; [w.textChanged.connect(self.update_preview) for w in widgets]; [c.toggled.connect(self.update_preview) for c in (self.show_uni_logo,self.show_fac_logo,self.show_author,self.show_university,self.show_career,self.show_year)]; self.update_preview()
    def apply_automatic_layout(self):
        self.title_font_size=30; self.text_font_size=15; self.margin_size=36
        self.layout_settings.update({'logo_uni_size':72,'logo_uni_x':0,'logo_uni_y':0,'logo_fac_size':58,'logo_fac_x':0,'logo_fac_y':0,'subtitle_x':0,'subtitle_y':0,'title_x':0,'title_y':0,'data_x':0,'data_y':0,'margin_top':36,'margin_bottom':36,'margin_left':42,'margin_right':42,'space_logo_subtitle':14,'space_subtitle_title':64,'space_title_data':34,'space_data_logo':12,'title_align':'center','subtitle_align':'center','data_align':'center','logo_align':'center','auto_fit':1})
    def design_in_word(self):
        try:
            data=self.values(); folder=DATA_DIR/'portadas_word'; folder.mkdir(parents=True,exist_ok=True); path=folder/(safe_name(data.get('title') or 'Portada')+'.docx'); self.create_cover_word(data,path); os.startfile(str(path))
        except Exception as exc:QMessageBox.warning(self,'Portada en Word',f'No se pudo crear la portada en Word.\n{exc}')
    def create_cover_word(self,data,path):
        doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
        def add(text,size,bold=False,after=6):
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(after);r=p.add_run(str(text));r.bold=bold;r.font.size=Pt(size);r.font.name=data.get('font_family') or 'Georgia'
        if data.get('show_university_logo',1) and data.get('university_logo') and Path(data['university_logo']).exists():p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(data['university_logo'],width=Inches(1.05))
        add((data.get('subtitle') or '').upper(),13,False,20);add((data.get('title') or 'TÍTULO DEL CUADERNO').upper(),24,True,28)
        for enabled,key,size in [(data.get('show_author',1),'author',14),(True,'faculty',13),(data.get('show_career',1),'career',13),(data.get('show_university',1),'university',13),(data.get('show_year',1),'year',13)]:
            if enabled and data.get(key):add(data[key],size)
        if data.get('show_faculty_logo',1) and data.get('faculty_logo') and Path(data['faculty_logo']).exists():p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(data['faculty_logo'],width=Inches(.85))
        doc.save(path)
    def edit_appearance(self):
        dialog=CoverAppearanceDialog(self.values(),self)
        if dialog.exec():
            values=dialog.values(); self.font_family=values['font_family']; self.font_color=values['font_color']; self.frame_style=values['frame_style']; self.apply_automatic_layout(); self.update_preview()
    def edit_assets(self):
        dialog=CoverAssetsDialog(dict(university_logo=self.uni_logo,faculty_logo=self.fac_logo,background_image=self.bg_image,color=self.palette.currentData(),background_position=self.background_position),self)
        if dialog.exec():
            values=dialog.values(); self.uni_logo=values['university_logo']; self.fac_logo=values['faculty_logo']; self.bg_image=values['background_image']; self.background_position=values['background_position']; self.palette.setCurrentIndex(max(0,self.palette.findData(values['color']))); self.update_preview()
    def pick_image(self,kind):
        path,_=QFileDialog.getOpenFileName(self,'Seleccionar imagen','','Imágenes (*.png *.jpg *.jpeg *.bmp *.webp);;Todos (*.*)')
        if not path:return
        if kind=='uni': self.uni_logo=path; self.uni_path.setText(path)
        elif kind=='fac': self.fac_logo=path; self.fac_path.setText(path)
        else:self.bg_image=path; self.bg_path.setText(path)
        self.update_preview()
    def logo_html(self,path,visible):
        if not visible or not path or not Path(path).exists():return ''
        return f"<img src='{QUrl.fromLocalFile(path).toString()}' width='54'><br>"
    def update_preview(self):
        try:
            data=self.values(); path=DATA_DIR/'previews'/'editor_portada.pdf'; create_cover_pdf_file(data,path); self.preview_pdf_doc.close(); self.preview_pdf_doc.load(str(path)); self.preview.setDocument(self.preview_pdf_doc)
        except Exception:
            pass
    def values(self):
        return dict(title=self.title.text().strip(),subtitle=self.subtitle.text().strip(),author=self.author.text().strip(),university=self.university.text().strip(),faculty=self.faculty.text().strip(),career=self.career.text().strip(),year=self.year.text().strip(),color=self.palette.currentData() or '#146C82',university_logo=self.uni_logo,faculty_logo=self.fac_logo,background_image=self.bg_image,background_position=self.background_position,font_family=self.font_family,font_color=self.font_color,frame_style=self.frame_style,margin_size=self.margin_size,title_font_size=self.title_font_size,text_font_size=self.text_font_size,show_university_logo=int(self.show_uni_logo.isChecked()),show_faculty_logo=int(self.show_fac_logo.isChecked()),show_author=int(self.show_author.isChecked()),show_university=int(self.show_university.isChecked()),show_career=int(self.show_career.isChecked()),show_year=int(self.show_year.isChecked()),word_path=self.word_path)

class SubjectDialog(QDialog):
    LEGAL_ICONS=['⚖','🏛','📜','📚','🔨','🛡','🖋','📖','🏢','🗂','🔍','✅','🧾','📋','🕊','🌐','👥','🏦','🚔','🔐','💼','🎓','📑','🗃','🏠','🏫','🏬','🏭','🗄','🗒','📰','✒','🖊','🖇','📌','📍','🔖','🔎','🔏','🔓','🚨','👮','🧑‍⚖️','👨‍⚖️','👩‍⚖️','🤝','🏆','🧠','💡','💰','💳','🪙','⚔','📢','📣','🗣','✍','📝','📂','📁','🗳','🪪','🛂','🌍','🧭','⏳','📅','📊','📈','📉','🔔','📨','📬','🧩','🔗','⛓','🕵','🚪','🔑','🧱','📕','📗','📘','📙','📓','📔','📒','📃','📄','📇','🧑‍🎓','🧑‍💼','🏥','🏞','🏙','🌎','🌏','🧮','💻','🖥','🗝','🧷','📏','📐','🪶','🕯','🖼','🏷','📦','🧰','⚠','❗','❓','ℹ']
    def __init__(self,parent=None):
        super().__init__(parent); self.setWindowTitle('Nueva materia'); self.resize(560,420); self.selected='⚖'; lay=QVBoxLayout(self); form=QFormLayout(); self.name=QLineEdit(); form.addRow('Nombre de la materia',self.name); lay.addLayout(form); title=QLabel('Elige un icono para la materia'); title.setStyleSheet('font-size:16px;font-weight:700;color:#F8FAFC'); lay.addWidget(title); note=QLabel('Solo puede seleccionarse un icono. El borde dorado indica la selección.'); note.setStyleSheet('color:#CBD5E1'); lay.addWidget(note)
        scroll=QScrollArea(); scroll.setWidgetResizable(True); scroll.setMinimumHeight(350)
        icon_host=QWidget(); icon_host.setObjectName('iconHost'); grid=QGridLayout(icon_host); grid.setSpacing(9); grid.setContentsMargins(12,12,12,12)
        self.icon_group=QButtonGroup(self); self.icon_group.setExclusive(True); self.icon_buttons=[]
        for i,icon in enumerate(self.LEGAL_ICONS):
            b=QPushButton(icon); b.setObjectName('legalIcon'); b.setCheckable(True); b.setMinimumSize(62,58); b.setMaximumSize(70,64); b.setStyleSheet("QPushButton#legalIcon{font-size:30px;background:#17263A;color:#F8FAFC;border:1px solid #3B4C61;border-radius:10px;padding:4px} QPushButton#legalIcon:hover{background:#22344B;border:1px solid #D4AF37} QPushButton#legalIcon:checked{background:#2A3D55;border:3px solid #D4AF37;color:#FFFFFF}")
            self.icon_group.addButton(b); self.icon_buttons.append(b); b.clicked.connect(lambda checked,x=icon:self.choose(x) if checked else None); grid.addWidget(b,i//9,i%9)
            if i==0: b.setChecked(True)
        scroll.setWidget(icon_host); lay.addWidget(scroll,1)
        self.preview=QLabel('Seleccionado: ⚖'); self.preview.setStyleSheet('color:#F8FAFC;font-weight:700'); lay.addWidget(self.preview)
        box=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel); box.accepted.connect(self.accept); box.rejected.connect(self.reject); lay.addWidget(box)
    def choose(self,icon): self.selected=icon; self.preview.setText('Seleccionado: '+icon)
    def download_icon(self,url):
        try:
            folder=DATA_DIR/'iconos'; folder.mkdir(exist_ok=True); path=folder/f'icono_{len(list(folder.glob("*")))+1}.png'; req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0'}); path.write_bytes(urllib.request.urlopen(req,timeout=20).read()); self.selected=str(path); self.preview.setText('Icono descargado: '+path.name)
        except Exception as e: QMessageBox.warning(self,'Icono',f'No se pudo descargar.\n{e}')
    def values(self): return self.name.text().strip(),self.selected

class ImageSearchDialog(QDialog):
    image_selected = Signal(str)
    def __init__(self,query,parent=None):
        super().__init__(parent); self.setWindowTitle('Buscar e insertar imagen'); self.resize(1050,720); lay=QVBoxLayout(self)
        bar=QHBoxLayout(); self.query=QLineEdit(query); go=QPushButton('Buscar'); go.setObjectName('primary'); go.clicked.connect(self.search); bar.addWidget(self.query,1); bar.addWidget(go); lay.addLayout(bar)
        info=QLabel('Haz clic sobre una imagen para descargarla e insertarla en la hoja.'); info.setObjectName('muted'); lay.addWidget(info)
        self.web=QWebEngineView(); self.web.loadFinished.connect(self.inject_picker); self.web.titleChanged.connect(self.title_changed); lay.addWidget(self.web,1); self.query.setFocus()
    def search(self):
        term=self.query.text().strip();
        if not term:return
        q=QUrl.toPercentEncoding(term).data().decode(); self.web.setUrl(QUrl('https://www.bing.com/images/search?q='+q))
    def inject_picker(self,ok=True):
        if not ok:return
        js="""(()=>{if(window.__lexPicker)return;window.__lexPicker=true;document.addEventListener('click',e=>{let i=e.target.closest('img');if(!i)return;e.preventDefault();e.stopPropagation();let u=i.currentSrc||i.src;if(u){document.title='LEXIMG:'+encodeURIComponent(u);}},true);})();"""
        self.web.page().runJavaScript(js)
    def title_changed(self,title):
        if not title.startswith('LEXIMG:'):return
        from urllib.parse import unquote
        url=unquote(title[7:]);
        if url.startswith('http'): self.image_selected.emit(url); self.accept()

class NavButton(QPushButton):
    def __init__(self,text): super().__init__(text); self.setCheckable(True); self.setCursor(Qt.PointingHandCursor)

class CoverPreview(QLabel):
    def __init__(self,parent=None):
        super().__init__(parent); self.cover_color='#146C82'; self.cover_pixmap=QPixmap(); self.cover_position='center'; self.frame_style='Doble dorado'; self.setAlignment(Qt.AlignCenter); self.setWordWrap(True); self.setMinimumSize(306,396); self.setMaximumSize(306,396); self.setAutoFillBackground(False); self.setAttribute(Qt.WA_TranslucentBackground,True); self.setStyleSheet('background:transparent;border:0')
    def set_cover_background(self,color,path='',position='center'):
        self.cover_color=color or '#146C82'; self.cover_pixmap=QPixmap(path) if path and Path(path).exists() else QPixmap(); self.cover_position=position or 'center'; self.update()
    def set_frame_style(self,style): self.frame_style=style or 'Doble dorado'; self.update()
    def paintEvent(self,event):
        painter=QPainter(self); rect=self.rect(); painter.fillRect(rect,QColor(self.cover_color))
        if not self.cover_pixmap.isNull() and rect.width()>0 and rect.height()>0:
            image=self.cover_pixmap.scaled(rect.size(),Qt.KeepAspectRatioByExpanding,Qt.SmoothTransformation); x=(rect.width()-image.width())//2; y=(rect.height()-image.height())//2
            if self.cover_position=='top': y=0
            elif self.cover_position=='bottom': y=rect.height()-image.height()
            elif self.cover_position=='left': x=0
            elif self.cover_position=='right': x=rect.width()-image.width()
            painter.drawPixmap(x,y,image)
        painter.end(); super().paintEvent(event); painter=QPainter(self); painter.setCompositionMode(QPainter.CompositionMode_SourceOver); paint_cover_frame(painter,self.rect().adjusted(2,2,-2,-2),self.frame_style); painter.end()

def safe_name(text):
    return ''.join(c if c.isalnum() or c in ' -_()' else '_' for c in text).strip() or 'Documento'

def word_available():
    try:
        import win32com.client
        app=win32com.client.DispatchEx('Word.Application'); app.Quit(); return True
    except Exception: return False

class BookView(QScrollArea):
    def __init__(self,parent=None):
        super().__init__(parent); self.setWidgetResizable(True); self.page=0; self.pdf=None; self.zoom=0.75; self.host=QWidget(); self.layout=QHBoxLayout(self.host); self.layout.setAlignment(Qt.AlignTop|Qt.AlignHCenter); self.setWidget(self.host)
    def clear_pages(self):
        while self.layout.count():
            item=self.layout.takeAt(0); w=item.widget()
            if w:w.deleteLater()
    def set_document(self,pdf): self.pdf=pdf; self.page=0; self.render_pages()
    def render_pages(self):
        self.clear_pages()
        if not self.pdf or self.pdf.pageCount()<1:return
        for idx in [self.page,self.page+1]:
            if idx>=self.pdf.pageCount():continue
            size=self.pdf.pagePointSize(idx); target=QSize(max(420,int(size.width()*self.zoom)),max(590,int(size.height()*self.zoom))); image=self.pdf.render(idx,target); label=QLabel(); label.setPixmap(QPixmap.fromImage(image)); label.setAlignment(Qt.AlignTop); label.setStyleSheet('background:white;border:1px solid #BFC8D2;padding:3px'); self.layout.addWidget(label)
    def next_spread(self):
        if self.pdf and self.page+2<self.pdf.pageCount(): self.page+=2; self.render_pages()
    def previous_spread(self):
        if self.page>0:self.page=max(0,self.page-2); self.render_pages()
    def zoom_in(self): self.zoom=min(1.6,self.zoom+0.1); self.render_pages()
    def zoom_out(self): self.zoom=max(0.35,self.zoom-0.1); self.render_pages()

class ReadingWindow(QMainWindow):
    def __init__(self,pdf_doc,title,parent=None):
        super().__init__(parent); self.pdf_doc=pdf_doc; self.setWindowTitle('Modo libro · '+title)
        self.setStyleSheet("QMainWindow,QWidget{background:#0B1220;color:#F8FAFC;font:14px 'Segoe UI'} QPushButton{background:#1F2937;color:#F8FAFC;border:1px solid #475569;border-radius:8px;padding:9px 14px} QPushButton:hover{background:#D4AF37;color:#111827}")
        root=QWidget(); self.setCentralWidget(root); v=QVBoxLayout(root); bar=QHBoxLayout(); heading=QLabel(title); heading.setStyleSheet('font:700 18px Georgia'); bar.addWidget(heading); bar.addStretch()
        for text,fn in [('‹ Anterior',lambda:self.book.previous_spread()),('Siguiente ›',lambda:self.book.next_spread()),('− Zoom',lambda:self.book.zoom_out()),('+ Zoom',lambda:self.book.zoom_in()),('Cerrar lectura',self.close)]: b=QPushButton(text); b.clicked.connect(fn); bar.addWidget(b)
        v.addLayout(bar); self.book=BookView(); self.book.set_document(pdf_doc); v.addWidget(self.book,1); self.showMaximized()

class Window(QMainWindow):
    def __init__(self):
        super().__init__(); self.store=Store(); self.cfg=load_cfg(); self.subject_id=None; self.unit_id=None; self.card_index=0
        self.setMinimumSize(1024,650); self.resize(1366,768); self.build_ui(); self.apply_style(); self.show_dashboard()
    def apply_style(self):
        a=self.cfg['accent']; dark=self.cfg['theme']=='dark'
        bg='#08111F' if dark else '#EEF2F6'; panel='#0E1B2B' if dark else '#FFFFFF'; panel2='#14263B' if dark else '#F6F8FB'; text='#F1F5F9' if dark else '#172334'; muted='#91A3B7' if dark else '#637387'; line='#21364D' if dark else '#DAE2EA'
        style = """
        * {font-family:'Segoe UI';font-size:13px} QMainWindow,QWidget#root {background:%s;color:%s} QFrame#sidebar,QFrame#card,QFrame#topbar {background:%s;border:1px solid %s;border-radius:14px}
        QLabel#brand {font:700 19px Georgia;color:%s} QLabel#title {font:700 25px Georgia;color:%s} QLabel#heading {font:700 19px Georgia;color:%s} QLabel#muted {color:%s} QLabel#accent {color:%s;font-weight:700}
        QPushButton {background:transparent;color:%s;border:0;border-radius:9px;padding:9px 12px;text-align:left;font-weight:600;min-height:20px} QPushButton:hover,QPushButton:checked {background:%s22;color:%s} QPushButton#primary {background:%s;color:#111827;text-align:center} QPushButton#outline {border:1px solid %s;color:%s;text-align:center} QPushButton#danger {border:1px solid #D96872;color:#D96872;text-align:center}
        QLineEdit,QTextEdit,QComboBox,QListWidget {background:%s;color:%s;border:1px solid %s;border-radius:9px;padding:8px;selection-background-color:%s;selection-color:#111827} QLineEdit:focus,QTextEdit:focus,QListWidget:focus {border:1px solid %s} QListWidget::item {padding:9px;border-radius:7px} QListWidget::item:selected {background:%s28;color:%s} QCheckBox {spacing:8px} QProgressBar {background:%s;border:0;border-radius:4px;height:7px} QProgressBar::chunk {background:%s;border-radius:4px} QScrollArea {border:0;background:transparent} QTabWidget::pane {border:1px solid %s;background:%s;border-radius:9px} QTabBar::tab {background:%s;color:%s;padding:8px 14px;border-top-left-radius:7px;border-top-right-radius:7px} QTabBar::tab:selected {background:%s;color:%s} QMenu {background:%s;color:%s;border:1px solid %s} QMenu::item:selected {background:%s;color:#111827} QDialog,QMessageBox {background:%s;color:%s} QDialog QLabel,QMessageBox QLabel,QFormLayout QLabel {color:#F1F5F9;font-size:14px;font-weight:600;font-family:'Segoe UI'} QComboBox QAbstractItemView {background:%s;color:%s;selection-background-color:%s;selection-color:#111827;border:1px solid %s} QCheckBox {color:%s}
        """ % (bg,text,panel,line,text,text,text,muted,a,muted,a,a,a,a,a,panel2,text,line,a,a,a,a,line,a, line,panel,panel2,muted,a,text,panel,text,line,a, panel,text,panel2,text,a,line,text)
        self.setStyleSheet(style)
        self.setWindowTitle(self.cfg['app_name'])
        self.brand.setText(f"⚖  {self.cfg['app_name']}")
        self.identity.setText(f"{self.cfg['user_name']}\n{self.cfg['career']}")
    def build_ui(self):
        root=QWidget(objectName='root'); self.setCentralWidget(root); outer=QHBoxLayout(root); outer.setContentsMargins(14,14,14,14); outer.setSpacing(14)
        side=QFrame(objectName='sidebar'); side.setMinimumWidth(205); side.setMaximumWidth(245); sv=QVBoxLayout(side); sv.setContentsMargins(15,18,15,15); sv.setSpacing(5)
        self.brand=QLabel(); self.brand.setObjectName('brand'); self.brand.setWordWrap(True); sv.addWidget(self.brand); self.identity=QLabel(); self.identity.setObjectName('muted'); sv.addWidget(self.identity); sv.addSpacing(18)
        self.nav=[]
        for text,fn in [('🏠   Inicio',self.show_dashboard),('📘   Portada Principal',self.show_main_cover_page),('📑   Índice General',self.show_main_index_page),('📚   Cuadernos',self.show_subjects),('◆   Flashcards',self.show_cards),('📝   Modo examen',self.show_exam),('📖   Biblioteca',self.show_library),('⚙   Configuración',self.show_settings)]:
            b=NavButton(text); b.clicked.connect(fn); sv.addWidget(b); self.nav.append(b)
        sv.addStretch(); backup=QPushButton('⇩   Exportar respaldo'); backup.clicked.connect(self.backup); sv.addWidget(backup); outer.addWidget(side)
        body=QWidget(); bv=QVBoxLayout(body); bv.setContentsMargins(0,0,0,0); bv.setSpacing(12)
        top=QFrame(objectName='topbar'); th=QHBoxLayout(top); th.setContentsMargins(18,12,18,12); self.page_title=QLabel(); self.page_title.setObjectName('title'); th.addWidget(self.page_title); th.addStretch(); self.search=QLineEdit(); self.search.setPlaceholderText('Buscar en tus apuntes...'); self.search.setMaximumWidth(330); self.search.returnPressed.connect(self.search_notes); th.addWidget(self.search); bv.addWidget(top)
        self.stack=QStackedWidget(); bv.addWidget(self.stack,1); outer.addWidget(body,1)
    def set_active(self,index):
        for i,b in enumerate(self.nav): b.setChecked(i==index)
    def page(self,title,index):
        while self.stack.count(): w=self.stack.widget(0); self.stack.removeWidget(w); w.deleteLater()
        self.page_title.setText(title); self.set_active(index); self.search.setVisible(index!=1); w=QWidget(); self.stack.addWidget(w); return w
    def card(self): return QFrame(objectName='card')
    def show_dashboard(self):
        w=self.page(f"Bienvenido, {self.cfg['user_name']}",0); v=QVBoxLayout(w); v.setContentsMargins(2,2,2,2); v.setSpacing(12)
        hero=self.card(); hh=QHBoxLayout(hero); hh.setContentsMargins(24,22,24,22); txt=QVBoxLayout(); k=QLabel('CENTRO DE ESTUDIO JURÍDICO'); k.setObjectName('accent'); txt.addWidget(k); h=QLabel('Estudia con orden, claridad y propósito'); h.setObjectName('title'); h.setWordWrap(True); txt.addWidget(h); m=QLabel(f"{self.cfg['university']} · {self.cfg['career']}"); m.setObjectName('muted'); txt.addWidget(m); hh.addLayout(txt,1); icon=QLabel('⚖'); icon.setStyleSheet('font-size:64px'); hh.addWidget(icon); v.addWidget(hero)
        stats=QGridLayout(); stats.setSpacing(12); values=[('Materias',self.store.value('SELECT COUNT(*) FROM subjects')),('Unidades',self.store.value('SELECT COUNT(*) FROM units')),('Completadas',self.store.value('SELECT COUNT(*) FROM units WHERE completed=1')),('Flashcards',self.store.value('SELECT COUNT(*) FROM cards'))]
        for i,(label,val) in enumerate(values):
            c=self.card(); cv=QVBoxLayout(c); cv.setContentsMargins(16,14,16,14); l=QLabel(label); l.setObjectName('muted'); n=QLabel(str(val)); n.setObjectName('title'); cv.addWidget(l); cv.addWidget(n); stats.addWidget(c,0,i)
        v.addLayout(stats); recent=QLabel('Continuar estudiando'); recent.setObjectName('heading'); v.addWidget(recent); grid=QGridLayout(); grid.setSpacing(12)
        for i,s in enumerate(self.store.rows('SELECT * FROM subjects ORDER BY id LIMIT 3')): grid.addWidget(self.subject_card(s),0,i)
        v.addLayout(grid); v.addStretch()
    def subject_card(self,s):
        c=self.card(); q=QVBoxLayout(c); q.setContentsMargins(16,16,16,16); icon=QLabel(); icon_value=s['icon'] or '⚖';
        if Path(icon_value).exists(): px=QPixmap(icon_value).scaled(34,34,Qt.KeepAspectRatio,Qt.SmoothTransformation); icon.setPixmap(px)
        else: icon.setText(icon_value); icon.setStyleSheet('font-size:27px')
        q.addWidget(icon); t=QLabel(s['name']); t.setObjectName('heading'); t.setWordWrap(True); q.addWidget(t); total=self.store.value('SELECT COUNT(*) FROM units WHERE subject_id=?',(s['id'],)); done=self.store.value('SELECT COUNT(*) FROM units WHERE subject_id=? AND completed=1',(s['id'],)); sub=QLabel(f'{total} unidades · {round(done/total*100) if total else 0}%'); sub.setObjectName('muted'); q.addWidget(sub); p=QProgressBar(); p.setRange(0,max(1,total)); p.setValue(done); p.setTextVisible(False); q.addWidget(p); b=QPushButton('Abrir cuaderno'); b.setObjectName('primary'); b.clicked.connect(lambda _,sid=s['id']:self.open_subject(sid)); q.addWidget(b); return c
    def show_subjects(self):
        self.subject_id=None; w=self.page('Cuadernos',3); v=QVBoxLayout(w); top=QHBoxLayout(); h=QLabel('Tus materias'); h.setObjectName('heading'); top.addWidget(h); top.addStretch(); b=QPushButton('+ Nueva materia'); b.setObjectName('primary'); b.clicked.connect(self.add_subject); top.addWidget(b); v.addLayout(top); grid=QGridLayout(); grid.setSpacing(12)
        for i,s in enumerate(self.store.rows('SELECT * FROM subjects ORDER BY id')): grid.addWidget(self.subject_card(s),i//3,i%3)
        v.addLayout(grid); v.addStretch()
    def open_subject(self,sid):
        self.subject_id=sid; subject=self.store.row('SELECT * FROM subjects WHERE id=?',(sid,)); w=self.page(subject['name'],3); h=QHBoxLayout(w); h.setSpacing(12)
        left=self.card(); left.setMinimumWidth(245); left.setMaximumWidth(310); lv=QVBoxLayout(left); main_cover=QPushButton('📘 Portada'); main_cover.setObjectName('outline'); main_cover.clicked.connect(self.edit_main_cover); lv.addWidget(main_cover); general_index=QPushButton('📑 Índice'); general_index.setObjectName('outline'); general_index.clicked.connect(self.show_general_index); lv.addWidget(general_index); line=QFrame(); line.setFrameShape(QFrame.HLine); line.setStyleSheet('color:#334155'); lv.addWidget(line); add=QPushButton('+ Nueva unidad'); add.setObjectName('primary'); add.clicked.connect(self.add_unit); lv.addWidget(add); self.unit_list=QListWidget(); self.unit_list.itemSelectionChanged.connect(self.load_unit); self.unit_list.setContextMenuPolicy(Qt.CustomContextMenu); self.unit_list.customContextMenuRequested.connect(self.show_unit_context_menu); lv.addWidget(self.unit_list,1); delete=QPushButton('Eliminar materia'); delete.setObjectName('danger'); delete.clicked.connect(self.delete_subject); lv.addWidget(delete); h.addWidget(left)
        viewer=self.card(); vv=QVBoxLayout(viewer); top=QHBoxLayout(); self.doc_title=QLabel('Selecciona una unidad'); self.doc_title.setObjectName('heading'); top.addWidget(self.doc_title); top.addStretch(); edit=QPushButton('✎ Editar en Microsoft Word'); edit.setObjectName('primary'); edit.clicked.connect(self.edit_in_word); top.addWidget(edit); refresh=QPushButton('↻ Actualizar vista'); refresh.setObjectName('outline'); refresh.clicked.connect(self.refresh_preview); top.addWidget(refresh); vv.addLayout(top)
        controls=QHBoxLayout(); self.single_btn=QPushButton('Visualización'); self.single_btn.setCheckable(True); self.single_btn.setChecked(True); self.single_btn.clicked.connect(lambda:self.set_read_mode(0)); controls.addWidget(self.single_btn); self.book_btn=QPushButton('Modo libro'); self.book_btn.setCheckable(True); self.book_btn.clicked.connect(self.open_book_window); controls.addWidget(self.book_btn); controls.addStretch(); vv.addLayout(controls)
        self.pdf_doc=QPdfDocument(self); self.pdf_view=QPdfView(); self.pdf_view.setDocument(self.pdf_doc); self.pdf_view.setPageMode(QPdfView.PageMode.SinglePage); self.pdf_view.setZoomMode(QPdfView.ZoomMode.FitInView); self.book_view=BookView(); self.viewer_stack=QStackedWidget(); self.viewer_stack.addWidget(self.pdf_view); self.viewer_stack.addWidget(self.book_view); vv.addWidget(self.viewer_stack,1); self.status=QLabel('El documento se mostrará aquí. Word se utilizará únicamente para editar.'); self.status.setObjectName('muted'); vv.addWidget(self.status); h.addWidget(viewer,1); self.reload_units()
    def first_subject_id(self):
        row=self.store.row('SELECT id FROM subjects ORDER BY id LIMIT 1')
        return row['id'] if row else None
    def subject_for_editorial(self,sid=None):
        sid=sid or self.subject_id or self.first_subject_id()
        return sid,self.store.row('SELECT * FROM subjects WHERE id=?',(sid,)) if sid else None
    def main_cover_data(self,sid=None):
        sid,subject=self.subject_for_editorial(sid); row=self.store.row('SELECT * FROM covers WHERE subject_id=?',(sid,)) if sid else None
        if row:return dict(row)
        return dict(subject_id=sid,title=subject['name'] if subject else '',subtitle='Cuaderno de estudio',author=self.cfg.get('user_name',''),university=self.cfg.get('university',''),faculty='',career=self.cfg.get('career',''),year=str(time.localtime().tm_year),color='#146C82',icon=subject['icon'] if subject else '⚖',university_logo='',faculty_logo='',background_image='',show_university_logo=1,show_faculty_logo=1,show_author=1,show_university=1,show_career=1,show_year=1,coverai_path='')
    def copy_cover_asset(self,path,kind,sid=None):
        sid=sid or self.subject_id or self.first_subject_id()
        if not sid or not path or not Path(path).exists():return ''
        folder=DATA_DIR/'portadas'/str(sid); folder.mkdir(parents=True,exist_ok=True); target=folder/(kind+Path(path).suffix.lower())
        if Path(path).resolve()!=target.resolve(): shutil.copy2(path,target)
        return str(target)
    def edit_main_cover(self,sid=None):
        sid=sid or self.subject_id or self.first_subject_id()
        if not sid:return
        data=self.main_cover_data(sid);folder=DATA_DIR/'coverai';folder.mkdir(parents=True,exist_ok=True);project=Path(data.get('coverai_path') or folder/f'Portada_Materia_{sid}.coverai');pdf=project.with_suffix('.pdf');self.subject_cover_studio=CoverAIStudio(project,pdf,data,self)
        def saved(proj,pdf_file):
            data['coverai_path']=proj;keys=['title','subtitle','author','university','faculty','career','year','color','icon','university_logo','faculty_logo','background_image','show_university_logo','show_faculty_logo','show_author','show_university','show_career','show_year','coverai_path'];columns=','.join(['subject_id']+keys);marks=','.join(['?']*(len(keys)+1));updates=','.join(f'{k}=excluded.{k}' for k in keys);self.store.run(f'INSERT INTO covers({columns}) VALUES({marks}) ON CONFLICT(subject_id) DO UPDATE SET {updates}',(sid,*[data.get(k) for k in keys]))
        self.subject_cover_studio.saved.connect(saved);self.subject_cover_studio.show()
    def global_cover_data(self):
        row=self.store.row('SELECT * FROM global_cover WHERE id=1')
        if row:return dict(row)
        return dict(id=1,title='Mi Cuaderno Jurídico',subtitle='Colección de cuadernos de estudio',author=self.cfg.get('user_name',''),university=self.cfg.get('university',''),faculty='',career=self.cfg.get('career',''),year=str(time.localtime().tm_year),color='#146C82',icon='⚖',university_logo='',faculty_logo='',background_image='',background_position='center',font_family='Georgia',font_color='#FFFFFF',frame_style='Doble dorado',margin_size=36,title_font_size=30,text_font_size=16,show_university_logo=1,show_faculty_logo=1,show_author=1,show_university=1,show_career=1,show_year=1,coverai_path='')
    def save_global_cover(self,x):
        folder=DATA_DIR/'portadas'/'principal'; folder.mkdir(parents=True,exist_ok=True)
        def copy_asset(path,name):
            if not path or not Path(path).exists():return ''
            target=folder/(name+Path(path).suffix.lower())
            if Path(path).resolve()!=target.resolve():shutil.copy2(path,target)
            return str(target)
        x=dict(x)
        x['university_logo']=copy_asset(x.get('university_logo',''),'logo_universidad')
        x['faculty_logo']=copy_asset(x.get('faculty_logo',''),'logo_facultad')
        x['background_image']=copy_asset(x.get('background_image',''),'fondo')
        keys=['title','subtitle','author','university','faculty','career','year','color','icon','university_logo','faculty_logo','background_image','background_position','font_family','font_color','frame_style','margin_size','title_font_size','text_font_size','show_university_logo','show_faculty_logo','show_author','show_university','show_career','show_year','coverai_path']
        columns=','.join(['id']+keys); marks=','.join(['?']*(len(keys)+1)); updates=','.join(f'{k}=excluded.{k}' for k in keys)
        sql=f'INSERT INTO global_cover({columns}) VALUES({marks}) ON CONFLICT(id) DO UPDATE SET {updates}'
        self.store.run(sql,(1,*[x.get(k) for k in keys]))
    def cover_html(self,d,compact=False):
        return build_cover_html(d)
    def show_main_cover_page(self):
        w=self.page('Portada Principal',1); layout=QHBoxLayout(w)
        editor=self.card(); ev=QVBoxLayout(editor); h=QLabel('Portada del documento'); h.setObjectName('heading'); ev.addWidget(h)
        info=QLabel('Edita la portada directamente con CoverAI Studio V7. Los cambios guardados se mantienen al cambiar de sección o reiniciar.'); info.setWordWrap(True); info.setObjectName('muted'); ev.addWidget(info)
        edit=QPushButton('Editar portada principal'); edit.setObjectName('primary'); ev.addWidget(edit)
        view_pdf=QPushButton('📄 Ver PDF'); view_pdf.setObjectName('outline'); ev.addWidget(view_pdf); ev.addStretch(); editor.setMaximumWidth(330); layout.addWidget(editor)
        preview_card=self.card(); pv=QVBoxLayout(preview_card); title=QLabel('Vista previa exacta tamaño Carta'); title.setObjectName('heading'); pv.addWidget(title)
        preview_doc=QPdfDocument(self); preview=QPdfView(); preview.setDocument(preview_doc); preview.setPageMode(QPdfView.PageMode.SinglePage); preview.setZoomMode(QPdfView.ZoomMode.FitInView); pv.addWidget(preview,1); layout.addWidget(preview_card,1)
        def refresh():
            data=self.global_cover_data(); custom=data.get('coverai_path') or ''; custom_pdf=Path(custom).with_suffix('.pdf') if custom else None
            path=custom_pdf if custom_pdf and custom_pdf.exists() else DATA_DIR/'previews'/'portada_principal.pdf'
            if not (custom_pdf and custom_pdf.exists()):create_cover_pdf_file(data,path)
            preview_doc.close(); preview_doc.load(str(path)); preview.setDocument(preview_doc)
        edit.clicked.connect(lambda:self.open_global_cover_studio(refresh)); view_pdf.clicked.connect(self.open_main_cover_pdf); refresh()
    def open_global_cover_studio(self,refresh_callback=None):
        data=self.global_cover_data(); folder=DATA_DIR/'coverai'; folder.mkdir(parents=True,exist_ok=True)
        project=Path(data.get('coverai_path') or folder/'Portada_Principal.coverai'); pdf=project.with_suffix('.pdf')
        data['coverai_path']=str(project); self.save_global_cover(data)
        self.cover_studio=CoverAIStudio(project,pdf,data,self)
        def saved(proj,pdf_file):
            current=self.global_cover_data(); current['coverai_path']=str(proj); self.save_global_cover(current)
            if refresh_callback:refresh_callback()
        self.cover_studio.saved.connect(saved); self.cover_studio.show()
    def open_main_cover_pdf(self):
        try:
            data=self.global_cover_data(); custom=data.get('coverai_path') or ''; custom_pdf=Path(custom).with_suffix('.pdf') if custom else None
            path=custom_pdf if custom_pdf and custom_pdf.exists() else DATA_DIR/'previews'/'portada_principal.pdf'
            if not (custom_pdf and custom_pdf.exists()):create_cover_pdf_file(data,path)
            dialog=QDialog(self); dialog.setWindowTitle('Vista PDF - Portada principal'); dialog.resize(760,600); layout=QVBoxLayout(dialog); doc=QPdfDocument(dialog); doc.load(str(path)); view=QPdfView(); view.setDocument(doc); view.setPageMode(QPdfView.PageMode.SinglePage); view.setZoomMode(QPdfView.ZoomMode.FitInView); layout.addWidget(view); dialog._pdf_doc=doc; dialog.exec()
        except Exception as exc:QMessageBox.warning(self,self.cfg['app_name'],f'No se pudo abrir la portada PDF.\n{exc}')
    def document_headings(self,path):
        result=[]; seen=set()
        try:
            doc=Document(str(path))
            for paragraph in doc.paragraphs:
                text=' '.join(paragraph.text.split())
                if not text:continue
                style=(paragraph.style.name or '').lower() if paragraph.style else ''
                level=None
                if 'heading' in style or 'título' in style or 'titulo' in style:
                    nums=re.findall(r'\d+',style); level=min(int(nums[0]),3) if nums else 1
                else:
                    match=re.match(r'^(\d+(?:\.\d+){1,3})\s*[.\-:]?\s+(.+)$',text)
                    if match: level=min(match.group(1).count('.')+1,3)
                key=text.casefold()
                if level and key not in seen: result.append((level,text)); seen.add(key)
        except Exception:pass
        return result
    def index_entries(self,sid=None):
        sid=sid or self.subject_id or self.first_subject_id(); entries=[]
        for number,u in enumerate(self.store.rows('SELECT * FROM units WHERE subject_id=? ORDER BY id',(sid,)),1):
            title=u['title']; unit_title=title if title.lower().startswith('unidad') else f'Unidad {number}: {title}'; topics=[]; path=Path(u['docx_path']) if u['docx_path'] else None
            if path and path.exists():
                for level,text in self.document_headings(path):
                    if text.casefold()!=title.casefold() and text.casefold()!=unit_title.casefold():topics.append((level,text))
            entries.append((unit_title,topics))
        return entries
    def general_index_text(self,sid=None):
        lines=['ÍNDICE GENERAL','']
        for unit,topics in self.index_entries(sid):
            lines.append(unit.upper())
            for level,text in topics:lines.append('    '*(max(level-1,0)) + text)
            lines.append('')
        return '\n'.join(lines).rstrip()
    def general_index_html(self,sid=None):
        rows=[]
        for unit,topics in self.index_entries(sid):
            rows.append(f"<div class='unit'><span>{html.escape(unit.upper())}</span><span class='dots'></span><span>pág. —</span></div>")
            for level,text in topics:
                margin=18+(level-1)*22; rows.append(f"<div class='topic' style='margin-left:{margin}px'><span>{html.escape(text)}</span><span class='dots'></span><span>—</span></div>")
        body=''.join(rows) or "<p class='empty'>Todavía no hay unidades.</p>"
        return "<html><head><style>body{background:#d9e1e8;margin:0;padding:24px;font-family:'Segoe UI';color:#183247}.page{background:#fff;max-width:720px;min-height:900px;margin:auto;padding:55px 58px;box-shadow:0 5px 18px #8090a0;border-top:12px solid #146C82}.title{text-align:center;color:#146C82;font-family:Georgia;font-size:28px;font-weight:700;letter-spacing:2px;margin-bottom:35px}.unit,.topic{display:flex;align-items:flex-end;margin:13px 0}.unit{font-weight:700;font-size:16px;color:#146C82;margin-top:24px}.topic{font-size:14px}.dots{flex:1;border-bottom:2px dotted #9aa8b5;margin:0 8px 5px}.note{text-align:center;color:#718096;margin-top:45px;font-size:12px}.empty{text-align:center;color:#718096}</style></head><body><div class='page'><div class='title'>CONTENIDO</div>"+body+"<div class='note'>Los números de página se calcularán al crear el PDF final.</div></div></body></html>"
    def global_index_entries(self):
        result=[]
        for subject in self.store.rows('SELECT * FROM subjects ORDER BY id'):
            result.append((subject['name'],self.index_entries(subject['id'])))
        return result
    def global_index_text(self):
        lines=['ÍNDICE GENERAL','']
        for subject,units in self.global_index_entries():
            lines.append(subject.upper())
            for unit,topics in units:
                lines.append('    '+unit.upper())
                for level,text in topics:lines.append('        '+'    '*(max(level-1,0))+text)
            lines.append('')
        return '\n'.join(lines).rstrip()
    def global_index_html(self):
        blocks=[]
        for subject,units in self.global_index_entries():
            blocks.append(f"<div class='subject'>{html.escape(subject.upper())}</div>")
            for unit,topics in units:
                blocks.append(f"<div class='unit'><span>{html.escape(unit)}</span><span class='dots'></span><span>—</span></div>")
                for level,text in topics:
                    margin=28+(level-1)*18; blocks.append(f"<div class='topic' style='margin-left:{margin}px'><span>{html.escape(text)}</span><span class='dots'></span><span>—</span></div>")
        body=''.join(blocks) or "<p class='empty'>Todavía no hay materias.</p>"
        return "<html><head><style>body{background:white;margin:0;padding:0;font-family:'Segoe UI';color:#222}.page{background:white;max-width:760px;min-height:940px;margin:auto;padding:48px 54px}.heading{text-align:left;font-family:Georgia;font-size:26px;font-weight:700;margin-bottom:32px}.subject{font-weight:800;font-size:16px;margin:22px 0 9px}.unit,.topic{display:flex;align-items:flex-end;margin:7px 0;font-size:14px}.unit{font-weight:600}.dots{flex:1;border-bottom:1px dotted #777;margin:0 8px 4px}.note{text-align:center;color:#777;margin-top:42px;font-size:12px}</style></head><body><div class='page'><div class='heading'>Contenido</div>"+body+"<div class='note'>La numeración real se añadirá al generar el PDF final.</div></div></body></html>"
    def show_index_dialog(self,sid=None,title='Índice'):
        sid=sid or self.subject_id or self.first_subject_id()
        if not sid:return
        d=QDialog(self); d.setWindowTitle(title); d.resize(850,700); v=QVBoxLayout(d); h=QLabel(title); h.setObjectName('heading'); v.addWidget(h); info=QLabel('Vista editorial basada en las unidades y en los títulos y subtítulos detectados en Word.'); info.setWordWrap(True); info.setObjectName('muted'); v.addWidget(info); preview=QTextBrowser(); preview.setHtml(self.general_index_html(sid)); v.addWidget(preview,1); row=QHBoxLayout(); refresh=QPushButton('↻ Actualizar índice'); refresh.setObjectName('outline'); refresh.clicked.connect(lambda:preview.setHtml(self.general_index_html(sid))); row.addWidget(refresh); copy=QPushButton('Copiar índice'); copy.setObjectName('primary'); copy.clicked.connect(lambda:QApplication.clipboard().setText(self.general_index_text(sid))); row.addWidget(copy); close=QPushButton('Cerrar'); close.clicked.connect(d.accept); row.addWidget(close); v.addLayout(row); d.exec()
    def show_general_index(self):
        self.show_index_dialog(self.subject_id,'Índice de la materia')
    def show_main_index_page(self):
        w=self.page('Índice General',2); layout=QVBoxLayout(w); top=self.card(); tv=QHBoxLayout(top); label=QLabel('Índice general del documento completo'); label.setObjectName('heading'); tv.addWidget(label); tv.addStretch(); refresh=QPushButton('↻ Actualizar'); refresh.setObjectName('outline'); tv.addWidget(refresh); copy=QPushButton('Copiar índice'); copy.setObjectName('primary'); tv.addWidget(copy); layout.addWidget(top); page=self.card(); pv=QVBoxLayout(page); preview=QTextBrowser(); preview.setStyleSheet('QTextBrowser{background:white;color:#222;border:0}'); pv.addWidget(preview); layout.addWidget(page,1)
        refresh.clicked.connect(lambda:preview.setHtml(self.global_index_html())); copy.clicked.connect(lambda:QApplication.clipboard().setText(self.global_index_text())); preview.setHtml(self.global_index_html())
    def reload_units(self):
        self.unit_list.clear()
        for u in self.store.rows('SELECT * FROM units WHERE subject_id=? ORDER BY id',(self.subject_id,)):
            it=QListWidgetItem(('✓ ' if u['completed'] else '○ ')+u['title']); it.setData(Qt.UserRole,u['id']); self.unit_list.addItem(it)
        if self.unit_list.count(): self.unit_list.setCurrentRow(0)
    def unit_folder(self):
        folder=DATA_DIR/'documentos'/str(self.subject_id); folder.mkdir(parents=True,exist_ok=True); return folder
    def ensure_unit_docx(self,u):
        path=Path(u['docx_path']) if u['docx_path'] else self.unit_folder()/(safe_name(u['title'])+'.docx')
        if not path.exists():
            doc=Document(); doc.add_heading(u['title'],0); doc.add_paragraph('Escribe aquí tus apuntes, conceptos, artículos, ejemplos y conclusiones.'); doc.save(path)
        if str(path)!=u['docx_path']: self.store.run('UPDATE units SET docx_path=? WHERE id=?',(str(path),u['id']))
        return path
    def preview_files(self):
        if not getattr(self,'current_docx',None): return []
        stem=Path(self.current_docx).stem
        return sorted(Path(self.current_docx).parent.glob(stem+'.preview.*.pdf'),key=lambda x:x.stat().st_mtime,reverse=True)
    def latest_preview(self):
        files=self.preview_files()
        legacy=Path(self.current_docx).with_suffix('.preview.pdf') if getattr(self,'current_docx',None) else None
        if legacy and legacy.exists(): files.append(legacy)
        return max(files,key=lambda x:x.stat().st_mtime) if files else None
    def load_unit(self):
        it=self.unit_list.currentItem()
        if not it:return
        self.unit_id=it.data(Qt.UserRole); u=self.store.row('SELECT * FROM units WHERE id=?',(self.unit_id,)); self.current_docx=self.ensure_unit_docx(u); self.doc_title.setText(u['title'])
        cached=self.latest_preview()
        if cached:
            self.current_pdf=cached; self.load_cached_pdf(); self.status.setText('Vista en caché cargada. Pulsa Actualizar vista si modificaste Word.')
        else:
            self.current_pdf=None; self.pdf_doc.close(); self.status.setText('Aún no existe una vista previa. Pulsa Actualizar vista.')
    def document_is_open_in_word(self):
        try:
            import win32com.client
            word=win32com.client.GetActiveObject('Word.Application'); target=str(Path(self.current_docx).resolve()).lower()
            for doc in word.Documents:
                try:
                    if str(Path(doc.FullName).resolve()).lower()==target:return True
                except Exception:pass
        except Exception:pass
        return False
    def load_cached_pdf(self):
        if not getattr(self,'current_pdf',None) or not Path(self.current_pdf).exists():return
        self.pdf_doc.close(); self.pdf_doc.load(str(self.current_pdf)); self.pdf_view.setDocument(self.pdf_doc); self.book_view.set_document(self.pdf_doc)
    def convert_word_to_pdf(self,docx_path,pdf_path):
        word=None; doc=None
        try:
            import win32com.client
            word=win32com.client.DispatchEx('Word.Application'); word.Visible=False; word.DisplayAlerts=0
            doc=word.Documents.Open(str(Path(docx_path).resolve()),ReadOnly=True,AddToRecentFiles=False); doc.ExportAsFixedFormat(str(Path(pdf_path).resolve()),17); doc.Close(False); doc=None; word.Quit(); word=None
            return True,''
        except Exception as e:
            try:
                if doc:doc.Close(False)
                if word:word.Quit()
            except Exception:pass
            Path(pdf_path).unlink(missing_ok=True); return False,str(e)
    def refresh_preview(self):
        if not getattr(self,'current_docx',None):return
        if self.document_is_open_in_word():
            cached=self.latest_preview()
            if cached:self.current_pdf=cached; self.load_cached_pdf()
            self.status.setText('El documento sigue abierto en Word. Guarda y cierra Word antes de actualizar.'); QMessageBox.information(self,self.cfg['app_name'],'Guarda y cierra el documento en Word. Después pulsa Actualizar vista.'); return
        cached=self.latest_preview()
        try:
            if cached and cached.stat().st_mtime>=Path(self.current_docx).stat().st_mtime:
                self.current_pdf=cached; self.load_cached_pdf(); self.status.setText('La visualización ya está actualizada.'); return
        except Exception:pass
        stamp=time.strftime('%Y%m%d_%H%M%S')+'_'+str(int(time.time()*1000)%1000).zfill(3)
        pdf=Path(self.current_docx).with_name(Path(self.current_docx).stem+f'.preview.{stamp}.pdf')
        self.status.setText('Generando una nueva vista previa con Microsoft Word...'); QApplication.processEvents(); ok,err=self.convert_word_to_pdf(self.current_docx,pdf)
        if not ok:
            if cached:self.current_pdf=cached; self.load_cached_pdf()
            self.status.setText('No se pudo actualizar. Se conserva la última visualización.'); QMessageBox.warning(self,self.cfg['app_name'],f'No se pudo actualizar la visualización.\n{err}'); return
        self.current_pdf=pdf; self.load_cached_pdf(); self.status.setText(f'Vista actualizada: {time.strftime("%H:%M:%S")}')
        # Limpieza segura: conserva las 5 vistas más recientes y nunca toca la que está cargada.
        for old in self.preview_files()[5:]:
            try:
                if old!=self.current_pdf: old.unlink()
            except OSError: pass
    def edit_in_word(self):
        if not getattr(self,'current_docx',None):return
        try: os.startfile(str(self.current_docx)); self.status.setText('Documento abierto en Word. Guarda los cambios y luego pulsa Actualizar vista.')
        except Exception as e: QMessageBox.warning(self,self.cfg['app_name'],f'No se pudo abrir Microsoft Word.\n{e}')
    def open_book_window(self):
        self.book_btn.setChecked(False)
        if not getattr(self,'current_pdf',None) or self.pdf_doc.pageCount()<1:
            QMessageBox.information(self,self.cfg['app_name'],'Primero selecciona una unidad y actualiza la visualización.'); return
        self.reading_window=ReadingWindow(self.pdf_doc,self.doc_title.text(),self); self.reading_window.show()
    def set_read_mode(self,index):
        self.viewer_stack.setCurrentIndex(0); self.single_btn.setChecked(True); self.book_btn.setChecked(False)
    def save_unit(self): pass
    def add_subject(self):
        d=SubjectDialog(self)
        if d.exec():
            name,icon=d.values()
            if name: self.store.run('INSERT INTO subjects(name,icon) VALUES(?,?)',(name,icon or '⚖')); self.show_subjects()
    def show_unit_context_menu(self,pos):
        item=self.unit_list.itemAt(pos)
        if not item:return
        self.unit_list.setCurrentItem(item); menu=QMenu(self); edit_action=menu.addAction('✏ Editar'); delete_action=menu.addAction('🗑 Eliminar'); chosen=menu.exec(self.unit_list.mapToGlobal(pos))
        if chosen==edit_action:self.edit_unit()
        elif chosen==delete_action:self.delete_unit()
    def delete_unit(self):
        if not hasattr(self,'unit_list') or not self.unit_list.currentItem():return
        unit_id=self.unit_list.currentItem().data(Qt.UserRole); row=self.store.row('SELECT * FROM units WHERE id=?',(unit_id,))
        if not row:return
        if QMessageBox.question(self,self.cfg['app_name'],f"¿Eliminar la unidad '{row['title']}'?",QMessageBox.Yes|QMessageBox.No)!=QMessageBox.Yes:return
        self.store.run('DELETE FROM units WHERE id=?',(unit_id,)); self.unit_id=None; self.current_docx=None; self.current_pdf=None
        if hasattr(self,'pdf_view'):self.pdf_doc.close(); self.doc_title.setText('Selecciona una unidad'); self.status.setText('Unidad eliminada.')
        self.reload_units()
    def edit_unit(self):
        if not hasattr(self,'unit_list') or not self.unit_list.currentItem():
            QMessageBox.information(self,self.cfg['app_name'],'Selecciona una unidad para editarla.'); return
        unit_id=self.unit_list.currentItem().data(Qt.UserRole); row=self.store.row('SELECT * FROM units WHERE id=?',(unit_id,))
        if not row:return
        new_title,ok=QInputDialog.getText(self,'Editar unidad','Nuevo nombre de la unidad:',text=row['title'])
        new_title=new_title.strip()
        if not ok or not new_title or new_title==row['title']:return
        self.store.run('UPDATE units SET title=? WHERE id=?',(new_title,unit_id))
        path=Path(row['docx_path']) if row['docx_path'] else None
        if path and path.exists():
            try:
                doc=Document(str(path))
                if doc.paragraphs and doc.paragraphs[0].text.strip()==row['title']:
                    doc.paragraphs[0].text=new_title; doc.save(str(path))
            except Exception:
                pass
        self.reload_units()
        for i in range(self.unit_list.count()):
            if self.unit_list.item(i).data(Qt.UserRole)==unit_id:self.unit_list.setCurrentRow(i);break
        if hasattr(self,'doc_title'):self.doc_title.setText(new_title)
    def add_unit(self):
        d=TextDialog('Nueva unidad',[('title','Título',False)],self)
        if d.exec() and d.values()['title']: self.store.run('INSERT INTO units(subject_id,title) VALUES(?,?)',(self.subject_id,d.values()['title'])); self.reload_units()
    def delete_subject(self):
        if QMessageBox.question(self,self.cfg['app_name'],'¿Eliminar esta materia y todo su contenido?')==QMessageBox.Yes:
            self.store.run('DELETE FROM units WHERE subject_id=?',(self.subject_id,)); self.store.run('DELETE FROM cards WHERE subject_id=?',(self.subject_id,)); self.store.run('DELETE FROM subjects WHERE id=?',(self.subject_id,)); self.show_subjects()
    def show_cards(self):
        rows=self.store.rows('SELECT cards.*,subjects.name subject FROM cards LEFT JOIN subjects ON subjects.id=cards.subject_id ORDER BY cards.id'); w=self.page('Flashcards',4); v=QVBoxLayout(w); top=QHBoxLayout(); top.addStretch(); b=QPushButton('+ Nueva tarjeta'); b.setObjectName('primary'); b.clicked.connect(self.add_card); top.addWidget(b); v.addLayout(top)
        if not rows:
            empty=QLabel('Aún no hay tarjetas. Pulsa Nueva tarjeta para crear la primera.'); empty.setObjectName('muted'); empty.setAlignment(Qt.AlignCenter); v.addWidget(empty,1); return
        self.card_index%=len(rows); r=rows[self.card_index]; c=self.card(); q=QVBoxLayout(c); q.setContentsMargins(30,30,30,30); q.addStretch(); s=QLabel(r['subject'] or 'Sin materia'); s.setObjectName('accent'); s.setAlignment(Qt.AlignCenter); q.addWidget(s); f=QLabel(r['front']); f.setObjectName('title'); f.setAlignment(Qt.AlignCenter); f.setWordWrap(True); q.addWidget(f); ans=QLabel(r['back']); ans.setObjectName('muted'); ans.setAlignment(Qt.AlignCenter); ans.setWordWrap(True); ans.hide(); q.addWidget(ans); show=QPushButton('Mostrar respuesta'); show.setObjectName('outline'); show.clicked.connect(ans.show); q.addWidget(show); q.addStretch(); v.addWidget(c,1); nxt=QPushButton('Siguiente tarjeta →'); nxt.setObjectName('primary'); nxt.clicked.connect(lambda:self.next_card(len(rows))); v.addWidget(nxt)
    def next_card(self,n): self.card_index=(self.card_index+1)%n; self.show_cards()
    def add_card(self):
        d=TextDialog('Nueva flashcard',[('front','Pregunta',True),('back','Respuesta',True)],self)
        if d.exec():
            x=d.values(); s=self.store.row('SELECT id FROM subjects ORDER BY id LIMIT 1')
            if x['front'] and x['back']: self.store.run('INSERT INTO cards(subject_id,front,back) VALUES(?,?,?)',(s['id'] if s else None,x['front'],x['back'])); self.show_cards()
    def show_exam(self):
        w=self.page('Modo examen',5); v=QVBoxLayout(w); exam_rows=self.store.rows('SELECT * FROM cards ORDER BY RANDOM() LIMIT 5'); intro=QLabel('Pulsa una pregunta para revelar la respuesta.'); intro.setObjectName('muted'); v.addWidget(intro)
        if not exam_rows:
            empty=QLabel('No hay preguntas disponibles. Primero crea tarjetas en Flashcards.'); empty.setObjectName('muted'); empty.setAlignment(Qt.AlignCenter); v.addWidget(empty,1); return
        for i,r in enumerate(exam_rows):
            c=self.card(); q=QVBoxLayout(c); b=QPushButton(f"{i+1}. {r['front']}"); a=QLabel(r['back']); a.setObjectName('accent'); a.setWordWrap(True); a.hide(); b.clicked.connect(a.show); q.addWidget(b); q.addWidget(a); v.addWidget(c)
        v.addStretch()
    def show_library(self):
        w=self.page('Biblioteca',6); v=QVBoxLayout(w); add=QPushButton('+ Añadir documento'); add.setObjectName('primary'); add.clicked.connect(self.add_file); v.addWidget(add); lst=QListWidget()
        for r in self.store.rows('SELECT * FROM library ORDER BY id DESC'):
            it=QListWidgetItem(f"{r['type']}   {r['name']}\n{r['path']}"); it.setData(Qt.UserRole,r['path']); lst.addItem(it)
        lst.itemDoubleClicked.connect(lambda it: os.startfile(it.data(Qt.UserRole))); v.addWidget(lst,1)
    def add_file(self):
        p,_=QFileDialog.getOpenFileName(self,'Añadir documento','','Documentos (*.pdf *.docx *.pptx *.txt *.png *.jpg);;Todos (*.*)')
        if p:
            x=Path(p); self.store.run('INSERT INTO library(name,path,type) VALUES(?,?,?)',(x.name,str(x),x.suffix[1:].upper())); self.show_library()
    def show_settings(self):
        w=self.page('Configuración',7)
        outer=QHBoxLayout(w)
        form=self.card()
        form.setObjectName('settingsCard')
        f=QFormLayout(form)
        f.setContentsMargins(24,24,24,24)
        f.setHorizontalSpacing(18)
        f.setVerticalSpacing(14)

        self.cfg_name=QLineEdit(self.cfg['app_name'])
        self.cfg_user=QLineEdit(self.cfg['user_name'])
        self.cfg_uni=QLineEdit(self.cfg['university'])
        self.cfg_career=QLineEdit(self.cfg['career'])
        self.cfg_accent=QComboBox()
        self.color_options=[('Dorado jurídico','#D4AF37'),('Azul profesional','#4EA1D3'),('Violeta académico','#9B7BD6'),('Borgoña universitario','#B86B77'),('Verde esmeralda','#4FB286')]
        [self.cfg_accent.addItem(name,code) for name,code in self.color_options]
        idx=self.cfg_accent.findData(self.cfg['accent'])
        self.cfg_accent.setCurrentIndex(max(0,idx))
        self.cfg_theme=QComboBox()
        self.cfg_theme.addItems(['dark','light'])
        self.cfg_theme.setCurrentText(self.cfg['theme'])

        dark=self.cfg.get('theme','dark')=='dark'
        label_color='#F8FAFC' if dark else '#172334'
        input_bg='#14263B' if dark else '#FFFFFF'
        input_text='#F8FAFC' if dark else '#172334'
        input_border='#3B4C61' if dark else '#CBD5E1'
        label_style=f"color:{label_color};font-family:'Segoe UI';font-size:14px;font-weight:600;background:transparent;"
        input_style=f"background:{input_bg};color:{input_text};border:1px solid {input_border};border-radius:9px;padding:8px;font-family:'Segoe UI';font-size:14px;"

        for field in (self.cfg_name,self.cfg_user,self.cfg_uni,self.cfg_career,self.cfg_accent,self.cfg_theme):
            field.setStyleSheet(input_style)

        def settings_label(text):
            label=QLabel(text)
            label.setStyleSheet(label_style)
            label.setMinimumWidth(175)
            return label

        f.addRow(settings_label('Nombre de la aplicación'),self.cfg_name)
        f.addRow(settings_label('Tu nombre'),self.cfg_user)
        f.addRow(settings_label('Universidad'),self.cfg_uni)
        f.addRow(settings_label('Carrera'),self.cfg_career)
        f.addRow(settings_label('Color principal'),self.cfg_accent)
        f.addRow(settings_label('Tema'),self.cfg_theme)

        save=QPushButton('Guardar configuración')
        save.setObjectName('primary')
        save.clicked.connect(self.save_settings)
        f.addRow(QLabel(''),save)
        outer.addWidget(form,1)
        outer.addStretch(1)
    def save_settings(self):
        self.cfg.update(app_name=self.cfg_name.text().strip() or DEFAULT_CFG['app_name'],user_name=self.cfg_user.text().strip() or DEFAULT_CFG['user_name'],university=self.cfg_uni.text().strip(),career=self.cfg_career.text().strip(),accent=self.cfg_accent.currentData(),theme=self.cfg_theme.currentText()); save_cfg(self.cfg); self.apply_style(); self.show_dashboard(); QMessageBox.information(self,self.cfg['app_name'],'Configuración guardada.')
    def search_notes(self):
        q=self.search.text().strip()
        if not q:return
        w=self.page(f'Resultados: {q}',-1); v=QVBoxLayout(w); rows=self.store.rows('SELECT * FROM units WHERE title LIKE ? OR content LIKE ?',(f'%{q}%',f'%{q}%'))
        for r in rows:
            c=self.card(); cv=QVBoxLayout(c); h=QLabel(r['title']); h.setObjectName('heading'); cv.addWidget(h); t=QLabel(r['content'][:350]); t.setObjectName('muted'); t.setWordWrap(True); cv.addWidget(t); v.addWidget(c)
        if not rows: v.addWidget(QLabel('No se encontraron coincidencias.'))
        v.addStretch()
    def backup(self):
        p,_=QFileDialog.getSaveFileName(self,'Exportar respaldo','Cuaderno-Juridico-Respaldo.json','JSON (*.json)')
        if p:
            data={'config':self.cfg,**{t:[dict(x) for x in self.store.rows(f'SELECT * FROM {t}')] for t in ['subjects','units','cards','library']}}; Path(p).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8'); QMessageBox.information(self,self.cfg['app_name'],'Respaldo exportado.')

if __name__=='__main__':
    app=QApplication(sys.argv); app.setStyle('Fusion'); win=Window(); win.show(); sys.exit(app.exec())
